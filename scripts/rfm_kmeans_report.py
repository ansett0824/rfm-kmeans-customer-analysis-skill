import os
import re
import sys
import traceback
import pandas as pd

from sklearn.preprocessing import StandardScaler
from sklearn.cluster import KMeans

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# RFM 與 K-means 客群分析報表產生工具
# 支援：
# 1. 明細銷貨資料：客戶名稱 / 交易日期 / 金額
# 2. 年度客戶別資料：客戶 / 客戶簡稱 / 1月~12月 / 合計
# 3. 客戶資料統計：客戶代號 / 客戶簡稱 / 客戶全名
# ============================================================

INPUT_DIR = "input"
OUTPUT_DIR = "output"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "客群分析結果報表.xlsx")
ERROR_LOG_FILE = os.path.join(OUTPUT_DIR, "error_log.txt")

MONTH_COLUMNS = [f"{i}月" for i in range(1, 13)]

CUSTOMER_NAME_CANDIDATES = [
    "客戶名稱",
    "客戶簡稱",
    "客戶全名",
    "客戶",
    "客戶原始欄位",
]

CUSTOMER_CODE_CANDIDATES = [
    "客戶代碼",
    "客戶代號",
    "客戶",
]

DATE_CANDIDATES = [
    "銷貨日期",
    "交易日期",
    "日期",
    "單據日期",
]

AMOUNT_CANDIDATES = [
    "銷售金額",
    "金額",
    "淨銷售金額",
    "合計",
]

PRODUCT_CANDIDATES = [
    "品名",
    "產品名稱",
    "品號",
]

QUANTITY_CANDIDATES = [
    "數量",
    "銷售數量",
    "總購買數量",
]


# ============================================================
# 基礎工具
# ============================================================

def print_title():
    print("=" * 70)
    print("RFM 與 K-means 客群分析報表產生工具")
    print("=" * 70)


def ensure_folders():
    os.makedirs(INPUT_DIR, exist_ok=True)
    os.makedirs(OUTPUT_DIR, exist_ok=True)


def write_error_log(message):
    ensure_folders()
    with open(ERROR_LOG_FILE, "w", encoding="utf-8") as f:
        f.write(message)


def find_excel_files():
    if not os.path.exists(INPUT_DIR):
        return []

    files = []
    for file in os.listdir(INPUT_DIR):
        if file.endswith(".xlsx") and not file.startswith("~$"):
            files.append(os.path.join(INPUT_DIR, file))

    return files


def read_excel_safely(file_path, sheet_name=0, nrows=None):
    try:
        return pd.read_excel(file_path, sheet_name=sheet_name, nrows=nrows)
    except Exception:
        return None


def get_excel_sheets(file_path):
    try:
        return pd.ExcelFile(file_path).sheet_names
    except Exception:
        return []


def find_first_existing_column(columns, candidates):
    for c in candidates:
        if c in columns:
            return c
    return None


def normalize_customer_name(value):
    if pd.isna(value):
        return None

    text = str(value).strip()

    if text == "":
        return None

    # 處理像「現售(11001)」這種格式，先保留括號前名稱
    text = re.sub(r"\(\s*\d+\s*\)$", "", text).strip()

    return text


def to_number(series):
    return pd.to_numeric(series, errors="coerce")


def parse_roc_or_ad_date(value):
    """
    支援：
    - 115/03/02
    - 2026/03/02
    - 2026-03-02
    - Excel 日期
    """
    if pd.isna(value):
        return pd.NaT

    if isinstance(value, pd.Timestamp):
        return value

    text = str(value).strip()

    if text == "":
        return pd.NaT

    # 民國格式：115/03/02 或 115-03-02
    m = re.match(r"^(\d{2,3})[/-](\d{1,2})[/-](\d{1,2})$", text)
    if m:
        year = int(m.group(1))
        month = int(m.group(2))
        day = int(m.group(3))

        # 民國年轉西元年
        if year < 1911:
            year += 1911

        try:
            return pd.Timestamp(year=year, month=month, day=day)
        except Exception:
            return pd.NaT

    return pd.to_datetime(value, errors="coerce")


def get_year_from_filename(file_path):
    """
    從檔名抓年度：
    例如 115年度銷貨資料-客戶別.xlsx → 115 → 2026
    """
    filename = os.path.basename(file_path)
    m = re.search(r"(\d{2,3})\s*年度", filename)

    if m:
        year = int(m.group(1))
        if year < 1911:
            return year + 1911
        return year

    return None


def month_end_date(year, month):
    start = pd.Timestamp(year=year, month=month, day=1)
    return start + pd.offsets.MonthEnd(0)


# ============================================================
# 檔案偵測
# ============================================================

def detect_file_types(excel_files):
    annual_candidates = []
    detail_candidates = []
    customer_candidates = []

    for file in excel_files:
        sheets = get_excel_sheets(file)

        for sheet in sheets:
            df_head = read_excel_safely(file, sheet_name=sheet, nrows=5)

            if df_head is None:
                continue

            columns = list(df_head.columns)

            month_count = sum(1 for m in MONTH_COLUMNS if m in columns)
            has_customer_name = any(c in columns for c in CUSTOMER_NAME_CANDIDATES)
            has_customer_code = any(c in columns for c in CUSTOMER_CODE_CANDIDATES)
            has_date = any(c in columns for c in DATE_CANDIDATES)
            has_amount = any(c in columns for c in AMOUNT_CANDIDATES)

            # 年度客戶別月份彙總資料
            if month_count >= 2 and has_customer_name:
                annual_candidates.append({
                    "file": file,
                    "sheet": sheet,
                    "month_count": month_count
                })

            # 明細交易資料
            if has_customer_name and has_date and has_amount:
                detail_candidates.append({
                    "file": file,
                    "sheet": sheet
                })

            # 客戶資料統計
            customer_score = 0
            if "客戶簡稱" in columns:
                customer_score += 2
            if "客戶全名" in columns:
                customer_score += 2
            if "客戶代號" in columns or "客戶代碼" in columns:
                customer_score += 2
            if "聯絡電話" in columns or "營業地址" in columns or "送貨地址" in columns:
                customer_score += 1

            # 避免把銷貨明細誤判成客戶資料
            if customer_score >= 3 and not has_date:
                customer_candidates.append({
                    "file": file,
                    "sheet": sheet,
                    "score": customer_score
                })

    annual_file = None
    detail_file = None
    customer_file = None

    if annual_candidates:
        annual_candidates = sorted(
            annual_candidates,
            key=lambda x: x["month_count"],
            reverse=True
        )
        annual_file = annual_candidates[0]

    if detail_candidates:
        # 優先使用 sheet 名稱包含「明細」的
        detail_candidates = sorted(
            detail_candidates,
            key=lambda x: 1 if "明細" in str(x["sheet"]) else 0,
            reverse=True
        )
        detail_file = detail_candidates[0]

    if customer_candidates:
        customer_candidates = sorted(
            customer_candidates,
            key=lambda x: x["score"],
            reverse=True
        )
        customer_file = customer_candidates[0]

    return annual_file, detail_file, customer_file


# ============================================================
# 資料整理：年度客戶別資料
# ============================================================

def build_rfm_from_annual_file(file_info):
    file = file_info["file"]
    sheet = file_info["sheet"]

    df = pd.read_excel(file, sheet_name=sheet)

    name_col = None
    if "客戶簡稱" in df.columns:
        name_col = "客戶簡稱"
    elif "客戶名稱" in df.columns:
        name_col = "客戶名稱"
    elif "客戶" in df.columns:
        name_col = "客戶"

    code_col = None
    if "客戶" in df.columns:
        code_col = "客戶"
    elif "客戶代號" in df.columns:
        code_col = "客戶代號"
    elif "客戶代碼" in df.columns:
        code_col = "客戶代碼"

    if name_col is None:
        raise ValueError("年度客戶別資料找不到客戶名稱欄位。")

    year = get_year_from_filename(file)
    if year is None:
        # 若檔名無法判斷年度，預設用目前資料常見的 115 年，即 2026
        year = 2026

    usable_months = [m for m in MONTH_COLUMNS if m in df.columns]

    records = []

    for _, row in df.iterrows():
        customer_name = normalize_customer_name(row.get(name_col))

        if customer_name is None:
            continue

        customer_code = row.get(code_col) if code_col else None

        monthly_values = []
        last_purchase_date = pd.NaT
        frequency = 0
        monetary = 0

        for month_col in usable_months:
            amount = pd.to_numeric(row.get(month_col), errors="coerce")

            if pd.isna(amount):
                amount = 0

            month_num = int(month_col.replace("月", ""))

            if amount > 0:
                frequency += 1
                monetary += amount
                last_purchase_date = month_end_date(year, month_num)

            monthly_values.append(amount)

        if monetary <= 0:
            continue

        record = {
            "客戶名稱": customer_name,
            "客戶代碼": customer_code,
            "LastPurchaseDate": last_purchase_date,
            "Frequency": frequency,
            "Monetary": monetary,
            "資料來源": os.path.basename(file),
            "資料格式": "年度客戶別月份彙總"
        }

        for month_col in usable_months:
            record[month_col] = pd.to_numeric(row.get(month_col), errors="coerce")

        if "合計" in df.columns:
            record["原始合計"] = pd.to_numeric(row.get("合計"), errors="coerce")

        records.append(record)

    rfm_base = pd.DataFrame(records)

    if rfm_base.empty:
        raise ValueError("年度客戶別資料沒有可用的正向銷售資料。")

    analysis_date = rfm_base["LastPurchaseDate"].max()
    rfm_base["Recency"] = (analysis_date - rfm_base["LastPurchaseDate"]).dt.days

    negative_sales = pd.DataFrame()

    return rfm_base, negative_sales, analysis_date


# ============================================================
# 資料整理：明細交易資料
# ============================================================

def build_rfm_from_detail_file(file_info):
    file = file_info["file"]
    sheet = file_info["sheet"]

    df = pd.read_excel(file, sheet_name=sheet)

    columns = list(df.columns)

    customer_col = find_first_existing_column(columns, ["客戶名稱", "客戶簡稱", "客戶全名", "客戶原始欄位", "客戶"])
    date_col = find_first_existing_column(columns, ["交易日期", "銷貨日期", "日期", "單據日期"])
    amount_col = find_first_existing_column(columns, ["金額", "銷售金額", "淨銷售金額", "合計"])

    code_col = find_first_existing_column(columns, ["客戶代碼", "客戶代號", "客戶"])

    if customer_col is None or date_col is None or amount_col is None:
        raise ValueError("明細資料缺少客戶、日期或金額欄位。")

    work = df.copy()
    work["客戶名稱"] = work[customer_col].apply(normalize_customer_name)
    work["銷貨日期"] = work[date_col].apply(parse_roc_or_ad_date)
    work["銷售金額"] = to_number(work[amount_col])

    if code_col:
        work["客戶代碼"] = work[code_col]

    work = work.dropna(subset=["客戶名稱", "銷貨日期", "銷售金額"])

    positive_sales = work[work["銷售金額"] > 0].copy()
    negative_sales = work[work["銷售金額"] <= 0].copy()

    if positive_sales.empty:
        raise ValueError("明細資料沒有可用的正向銷售資料。")

    analysis_date = positive_sales["銷貨日期"].max()

    agg_dict = {
        "LastPurchaseDate": ("銷貨日期", "max"),
        "Frequency": ("銷貨日期", "count"),
        "Monetary": ("銷售金額", "sum")
    }

    rfm_base = positive_sales.groupby("客戶名稱").agg(**agg_dict).reset_index()

    if "客戶代碼" in positive_sales.columns:
        code_map = positive_sales.groupby("客戶名稱")["客戶代碼"].first().reset_index()
        rfm_base = rfm_base.merge(code_map, on="客戶名稱", how="left")

    rfm_base["Recency"] = (analysis_date - rfm_base["LastPurchaseDate"]).dt.days
    rfm_base["資料來源"] = os.path.basename(file)
    rfm_base["資料格式"] = "交易明細"

    return rfm_base, negative_sales, analysis_date


# ============================================================
# 客戶資料整理
# ============================================================

def load_customer_data(customer_info):
    if customer_info is None:
        return pd.DataFrame()

    file = customer_info["file"]
    sheet = customer_info["sheet"]

    df = pd.read_excel(file, sheet_name=sheet)

    if df.empty:
        return pd.DataFrame()

    result = df.copy()

    # 建立統一的客戶名稱欄位
    if "客戶簡稱" in result.columns:
        result["客戶名稱"] = result["客戶簡稱"].apply(normalize_customer_name)
    elif "客戶名稱" in result.columns:
        result["客戶名稱"] = result["客戶名稱"].apply(normalize_customer_name)
    elif "客戶全名" in result.columns:
        result["客戶名稱"] = result["客戶全名"].apply(normalize_customer_name)
    else:
        return pd.DataFrame()

    # 統一客戶代碼欄位
    if "客戶代號" in result.columns and "客戶代碼" not in result.columns:
        result["客戶代碼"] = result["客戶代號"]

    result = result.dropna(subset=["客戶名稱"])
    result = result.drop_duplicates(subset=["客戶名稱"], keep="first")

    useful_cols = [
        "客戶名稱",
        "客戶代碼",
        "客戶簡稱",
        "客戶全名",
        "客戶地區",
        "聯絡人",
        "聯絡電話",
        "聯絡人手機",
        "營業地址",
        "送貨地址",
        "業務員",
        "業務員姓名",
        "最近交易日",
        "客戶型態",
        "客戶等級",
        "客戶分級",
        "客戶行業別",
    ]

    keep_cols = [c for c in useful_cols if c in result.columns]

    return result[keep_cols].copy()


# ============================================================
# RFM 分數與分群
# ============================================================

def safe_qcut(series, q=5, higher_is_better=True):
    """
    將資料分成 1~5 分。
    higher_is_better=True 代表數值越大分數越高。
    higher_is_better=False 代表數值越小分數越高。
    """
    s = pd.to_numeric(series, errors="coerce").fillna(0)

    if not higher_is_better:
        s = -s

    try:
        return pd.qcut(
            s.rank(method="first"),
            q,
            labels=list(range(1, q + 1))
        ).astype(int)
    except Exception:
        try:
            return pd.cut(
                s.rank(method="first"),
                bins=q,
                labels=list(range(1, q + 1)),
                include_lowest=True
            ).astype(int)
        except Exception:
            return pd.Series([3] * len(s), index=s.index)


def add_rfm_scores_and_clusters(rfm_base):
    rfm = rfm_base.copy()

    rfm["R_Score"] = safe_qcut(rfm["Recency"], 5, higher_is_better=False)
    rfm["F_Score"] = safe_qcut(rfm["Frequency"], 5, higher_is_better=True)
    rfm["M_Score"] = safe_qcut(rfm["Monetary"], 5, higher_is_better=True)

    rfm["RFM_Total"] = rfm["R_Score"] + rfm["F_Score"] + rfm["M_Score"]

    customer_count = len(rfm)
    cluster_count = 3 if customer_count >= 3 else customer_count

    if cluster_count >= 2:
        features = rfm[["R_Score", "F_Score", "M_Score"]]

        scaler = StandardScaler()
        features_scaled = scaler.fit_transform(features)

        kmeans = KMeans(n_clusters=cluster_count, random_state=42, n_init=10)
        rfm["Cluster"] = kmeans.fit_predict(features_scaled)

        cluster_summary = rfm.groupby("Cluster")[[
            "R_Score",
            "F_Score",
            "M_Score",
            "RFM_Total",
            "Monetary",
            "Frequency",
            "Recency"
        ]].mean().reset_index()

        cluster_order = cluster_summary.sort_values(
            by="RFM_Total",
            ascending=False
        )["Cluster"].tolist()

        if cluster_count == 3:
            segment_map = {
                cluster_order[0]: "高價值客戶",
                cluster_order[1]: "一般客戶",
                cluster_order[2]: "可能流失客戶"
            }
        else:
            segment_map = {
                cluster_order[0]: "較高價值客戶",
                cluster_order[1]: "需追蹤客戶"
            }

        rfm["客戶類型"] = rfm["Cluster"].map(segment_map)

    else:
        rfm["Cluster"] = 0
        rfm["客戶類型"] = "單一客戶"
        cluster_summary = rfm.groupby("Cluster")[[
            "R_Score",
            "F_Score",
            "M_Score",
            "RFM_Total",
            "Monetary",
            "Frequency",
            "Recency"
        ]].mean().reset_index()

    return rfm, cluster_summary


# ============================================================
# 其他摘要分析
# ============================================================

def create_segment_summary(result_df):
    total_sales = result_df["Monetary"].sum()

    segment_summary = result_df.groupby("客戶類型").agg(
        客戶數=("客戶名稱", "count"),
        銷售金額=("Monetary", "sum"),
        平均RFM分數=("RFM_Total", "mean"),
        平均購買頻率=("Frequency", "mean"),
        平均最近購買間隔=("Recency", "mean")
    ).reset_index()

    if total_sales > 0:
        segment_summary["銷售占比"] = segment_summary["銷售金額"] / total_sales
    else:
        segment_summary["銷售占比"] = 0

    return segment_summary


def create_monthly_summary(rfm_base):
    month_cols = [c for c in MONTH_COLUMNS if c in rfm_base.columns]

    if not month_cols:
        return pd.DataFrame()

    rows = []
    for month in month_cols:
        amount = pd.to_numeric(rfm_base[month], errors="coerce").fillna(0).sum()
        customers = (pd.to_numeric(rfm_base[month], errors="coerce").fillna(0) > 0).sum()
        rows.append({
            "月份": month,
            "銷售金額": amount,
            "交易客戶數": customers
        })

    return pd.DataFrame(rows)


def create_top_customer_summary(result_df, n=20):
    cols = [
        "客戶名稱",
        "客戶類型",
        "Monetary",
        "Frequency",
        "Recency",
        "RFM_Total",
        "LastPurchaseDate"
    ]

    keep_cols = [c for c in cols if c in result_df.columns]

    return result_df.sort_values(by="Monetary", ascending=False)[keep_cols].head(n)


# ============================================================
# Excel 美化
# ============================================================

def format_excel(file_path):
    wb = load_workbook(file_path)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center

        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)

            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

            ws.column_dimensions[col_letter].width = min(max_length + 4, 38)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                header = ws.cell(row=1, column=cell.column).value
                header_text = str(header)

                if isinstance(cell.value, (int, float)):
                    if any(k in header_text for k in ["金額", "Monetary", "銷售"]):
                        cell.number_format = '#,##0'
                    elif "占比" in header_text or "%" in header_text:
                        cell.number_format = '0.00%'
                    elif "分數" in header_text or "Score" in header_text:
                        cell.number_format = '0.00'
                    else:
                        cell.number_format = '#,##0.00'

    wb.save(file_path)


# ============================================================
# 主程式
# ============================================================

def main():
    print_title()
    ensure_folders()

    excel_files = find_excel_files()

    if not excel_files:
        message = (
            "找不到 Excel 檔案。\n"
            "請將 Excel 檔案放入 input 資料夾後，再重新執行。\n"
        )
        print(message)
        write_error_log(message)
        input("按 Enter 結束...")
        sys.exit(1)

    print("已找到以下 Excel 檔案：")
    for f in excel_files:
        print(f"- {f}")

    annual_info, detail_info, customer_info = detect_file_types(excel_files)

    print("\n檔案偵測結果：")
    print(f"年度客戶別資料：{annual_info['file'] if annual_info else '未偵測到'}")
    print(f"交易明細資料：{detail_info['file'] if detail_info else '未偵測到'}")
    print(f"客戶資料統計：{customer_info['file'] if customer_info else '未偵測到'}")

    try:
        # 優先使用年度客戶別資料，因為它涵蓋較完整年度資料
        if annual_info:
            print("\n使用年度客戶別月份資料進行 RFM 分析。")
            rfm_base, negative_sales, analysis_date = build_rfm_from_annual_file(annual_info)
        elif detail_info:
            print("\n使用交易明細資料進行 RFM 分析。")
            rfm_base, negative_sales, analysis_date = build_rfm_from_detail_file(detail_info)
        else:
            raise ValueError(
                "找不到可用的銷貨資料。\n"
                "可接受格式：\n"
                "1. 年度客戶別資料：客戶 / 客戶簡稱 / 1月~12月 / 合計\n"
                "2. 明細資料：客戶名稱 / 交易日期 / 金額\n"
            )

        customer_df = load_customer_data(customer_info)

        rfm, cluster_summary = add_rfm_scores_and_clusters(rfm_base)

        if not customer_df.empty:
            result = rfm.merge(customer_df, on="客戶名稱", how="left", suffixes=("", "_客戶資料"))
        else:
            result = rfm.copy()

        result = result.sort_values(by="RFM_Total", ascending=False)

        total_customers = len(result)
        total_sales = result["Monetary"].sum()

        overview = pd.DataFrame({
            "項目": [
                "分析基準日",
                "分析資料格式",
                "總客戶數",
                "正向銷售總金額",
                "是否偵測到年度客戶別資料",
                "是否偵測到交易明細資料",
                "是否偵測到客戶資料統計",
                "輸出報表"
            ],
            "內容": [
                str(analysis_date.date()) if not pd.isna(analysis_date) else "無法判斷",
                rfm_base["資料格式"].iloc[0] if "資料格式" in rfm_base.columns else "未標示",
                total_customers,
                total_sales,
                "是" if annual_info else "否",
                "是" if detail_info else "否",
                "是" if customer_info else "否",
                OUTPUT_FILE
            ]
        })

        segment_summary = create_segment_summary(result)
        monthly_summary = create_monthly_summary(rfm_base)
        top_customers = create_top_customer_summary(result)

        high_value_df = result[result["客戶類型"].str.contains("高價值", na=False)].copy()
        follow_up_df = result[result["客戶類型"].str.contains("流失|追蹤", na=False)].copy()

        with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
            overview.to_excel(writer, sheet_name="報表總覽", index=False)
            result.to_excel(writer, sheet_name="客群分析結果", index=False)
            cluster_summary.to_excel(writer, sheet_name="分群摘要", index=False)
            segment_summary.to_excel(writer, sheet_name="客戶類型摘要", index=False)
            top_customers.to_excel(writer, sheet_name="前20大客戶", index=False)
            high_value_df.to_excel(writer, sheet_name="高價值客戶", index=False)
            follow_up_df.to_excel(writer, sheet_name="需追蹤客戶", index=False)

            if not monthly_summary.empty:
                monthly_summary.to_excel(writer, sheet_name="月份銷售摘要", index=False)

            if negative_sales is not None and not negative_sales.empty:
                negative_sales.to_excel(writer, sheet_name="負值或退貨資料", index=False)

        format_excel(OUTPUT_FILE)

        print("\n報表產生完成！")
        print(f"輸出位置：{OUTPUT_FILE}")
        print("\n請至 output 資料夾查看 Excel 報表。")

        # 成功時清空舊錯誤紀錄
        if os.path.exists(ERROR_LOG_FILE):
            os.remove(ERROR_LOG_FILE)

        input("\n按 Enter 結束...")

    except Exception as e:
        error_message = (
            "報表產生失敗。\n\n"
            f"錯誤原因：{str(e)}\n\n"
            "詳細錯誤：\n"
            f"{traceback.format_exc()}\n"
        )

        print("\n" + error_message)
        write_error_log(error_message)
        print(f"錯誤紀錄已輸出至：{ERROR_LOG_FILE}")
        input("\n按 Enter 結束...")
        sys.exit(1)


if __name__ == "__main__":
    main()
